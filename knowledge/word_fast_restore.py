"""Fast restore structured historical Word knowledge for Update mode."""

from __future__ import annotations

import json
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from docx import Document

from knowledge.adapter import build_knowledge_id, normalize_internal_knowledge_type
from knowledge.models import KnowledgeItem, normalize_question


FAST_RESTORE_STRATEGY_VERSION = "word-fast-restore-v1"
DEEP_RESTORE_REMAINDER_THRESHOLD = 2000
MIN_FAST_RESTORE_CONFIDENCE = 0.72

QUESTION_HEADERS = {"问题", "问", "question", "q", "faq", "用户问题"}
ANSWER_HEADERS = {"回答", "答", "answer", "a", "话术", "回复", "标准回答"}
CATEGORY_HEADERS = {"分类", "类别", "category", "模块"}
MODEL_HEADERS = {"车型", "model"}
TRIM_HEADERS = {"版本", "版型", "trim"}
BRAND_HEADERS = {"品牌", "brand"}


@dataclass(slots=True)
class FastRestoreResult:
    items: list[KnowledgeItem] = field(default_factory=list)
    remaining_text: str = ""
    document_text: str = ""
    total_chars: int = 0
    remaining_chars: int = 0
    consumed_blocks: int = 0
    table_items: int = 0
    qa_items: int = 0
    heading_items: int = 0
    json_items: int = 0
    failed_items: int = 0
    confidence: float = 0.0

    @property
    def remaining_ratio(self) -> float:
        if not self.total_chars:
            return 0.0
        return self.remaining_chars / self.total_chars


def get_source_name(file: Any) -> str:
    name = getattr(file, "name", None)
    if name:
        return Path(str(name)).name
    if isinstance(file, (str, Path)):
        return Path(file).name
    return "历史知识.docx"


def load_document(file: Any) -> Document:
    if hasattr(file, "seek"):
        file.seek(0)
    return Document(file)


def clean_text(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip())


def contains_json_marker(text: str) -> bool:
    return bool(re.search(r'["“”]?(?:问题|答案|回答|车型|车系|分类|版本)["“”]?\s*[:：]', str(text or "")))


def looks_like_json_text(text: str) -> bool:
    text = str(text or "").strip()
    return ("{" in text and "}" in text and contains_json_marker(text)) or (
        text.startswith("[") and contains_json_marker(text)
    )


def normalize_header(value: str) -> str:
    return clean_text(value).lower().replace("：", "").replace(":", "")


def header_index(headers: list[str], candidates: set[str]) -> int | None:
    normalized = [normalize_header(header) for header in headers]
    for index, header in enumerate(normalized):
        if header in candidates:
            return index
    for index, header in enumerate(normalized):
        if any(candidate in header for candidate in candidates):
            return index
    return None


def infer_category(question: str, answer: str, fallback: str = "") -> str:
    if fallback:
        return fallback

    text = f"{question} {answer}"
    category_keywords = [
        ("价格", ["价格", "售价", "指导价", "多少钱", "价位"]),
        ("金融", ["金融", "贷款", "免息", "分期", "月供"]),
        ("权益", ["权益", "优惠", "补贴", "活动", "赠送"]),
        ("续航", ["续航", "公里", "电耗"]),
        ("动力", ["动力", "电机", "发动机", "加速"]),
        ("空间", ["空间", "轴距", "后排", "座椅"]),
        ("智驾", ["智驾", "辅助驾驶", "泊车", "领航"]),
        ("安全", ["安全", "气囊", "碰撞"]),
        ("配置", ["配置", "功能", "屏幕", "车机"]),
    ]
    for category, keywords in category_keywords:
        if any(keyword in text for keyword in keywords):
            return category
    return "通用"


def make_item(
    *,
    source_file: str,
    index: int,
    question: str,
    answer: str,
    category: str = "",
    model: str = "",
    trim: str = "",
    brand: str = "",
    source_type: str,
) -> KnowledgeItem | None:
    question = clean_text(question)
    answer = clean_text(answer)
    if not question or not answer:
        return None
    if len(question) < 2 or len(answer) < 2:
        return None
    if not is_valid_question_answer(question, answer):
        return None

    category = infer_category(question, answer, clean_text(category))
    return KnowledgeItem(
        knowledge_id=build_knowledge_id("DOCX", source_file, index, question, answer, category),
        question=question,
        normalized_question=normalize_question(question),
        answer=answer,
        category=category,
        module=None,
        brand=clean_text(brand) or None,
        model=clean_text(model) or None,
        trim=clean_text(trim) or None,
        knowledge_type=normalize_internal_knowledge_type(None, category, None),
        answer_type="restored",
        need_confirm=False,
        fact_refs=[],
        source_files=[source_file],
        metadata={
            "source": "word_fast_restore",
            "source_type": source_type,
            "strategy_version": FAST_RESTORE_STRATEGY_VERSION,
        },
    )


def is_valid_question_answer(question: str, answer: str) -> bool:
    question = clean_text(question)
    answer = clean_text(answer)
    if looks_like_json_text(question) or looks_like_json_text(answer):
        return False
    if contains_json_marker(question) or contains_json_marker(answer):
        return False
    if any(mark in question for mark in ["{", "}", "[", "]"]):
        return False
    if any(mark in answer for mark in ["{", "}", "[", "]"]):
        return False
    if len(question) > 120:
        return False
    if len(answer) > 1500:
        return False
    if "是什么？" in question and ("答案" in question or "回答" in question):
        return False
    return True


def normalize_json_quotes(text: str) -> str:
    return (
        str(text or "")
        .replace("“", '"')
        .replace("”", '"')
        .replace("‘", "'")
        .replace("’", "'")
        .strip()
    )


def parse_json_values(text: str) -> list[Any]:
    text = normalize_json_quotes(text)
    values: list[Any] = []

    candidates = [text]
    if "}{" in text:
        candidates.append(f"[{text.replace('}{', '},{')}]")

    for candidate in candidates:
        try:
            parsed = json.loads(candidate)
            if isinstance(parsed, list):
                values.extend(parsed)
            else:
                values.append(parsed)
            return values
        except Exception:
            pass

    decoder = json.JSONDecoder()
    index = 0
    while index < len(text):
        brace_index = text.find("{", index)
        if brace_index < 0:
            break
        try:
            parsed, end_index = decoder.raw_decode(text[brace_index:])
            values.append(parsed)
            index = brace_index + end_index
        except Exception:
            index = brace_index + 1

    return values


def dict_value(data: dict[str, Any], *keys: str) -> str:
    for key in keys:
        if key in data and data[key] is not None:
            return clean_text(data[key])
    return ""


def json_dict_to_item(data: dict[str, Any], source_file: str, index: int) -> KnowledgeItem | None:
    question = dict_value(data, "问题", "question", "Q", "q", "用户问题")
    answer = dict_value(data, "答案", "回答", "answer", "A", "a", "标准回答")
    model = dict_value(data, "车型", "车系", "model")
    trim = dict_value(data, "版本", "版型", "trim")
    brand = dict_value(data, "品牌", "brand")
    category = dict_value(data, "分类", "类别", "category", "模块")
    return make_item(
        source_file=source_file,
        index=index,
        question=question,
        answer=answer,
        category=category,
        model=model,
        trim=trim,
        brand=brand,
        source_type="json",
    )


def restore_json_blocks(lines: list[str], source_file: str) -> tuple[list[KnowledgeItem], set[int], int]:
    items: list[KnowledgeItem] = []
    consumed: set[int] = set()
    failed = 0
    index = 30000

    for line_index, line in enumerate(lines):
        if not looks_like_json_text(line):
            if contains_json_marker(line):
                consumed.add(line_index)
                failed += 1
            continue
        consumed.add(line_index)
        values = parse_json_values(line)
        if not values:
            failed += 1
            continue
        for value in values:
            if isinstance(value, list):
                dicts = [item for item in value if isinstance(item, dict)]
            elif isinstance(value, dict):
                dicts = [value]
            else:
                dicts = []
            for data in dicts:
                item = json_dict_to_item(data, source_file, index)
                if item:
                    items.append(item)
                    index += 1
                else:
                    failed += 1

    return items, consumed, failed


def restore_tables(doc: Document, source_file: str) -> tuple[list[KnowledgeItem], list[str]]:
    items: list[KnowledgeItem] = []
    consumed_texts: list[str] = []
    index = 1

    for table in doc.tables:
        rows = [
            [clean_text(cell.text) for cell in row.cells]
            for row in table.rows
        ]
        rows = [row for row in rows if any(row)]
        if len(rows) < 2:
            continue

        headers = rows[0]
        question_idx = header_index(headers, QUESTION_HEADERS)
        answer_idx = header_index(headers, ANSWER_HEADERS)
        if question_idx is None or answer_idx is None:
            continue

        category_idx = header_index(headers, CATEGORY_HEADERS)
        model_idx = header_index(headers, MODEL_HEADERS)
        trim_idx = header_index(headers, TRIM_HEADERS)
        brand_idx = header_index(headers, BRAND_HEADERS)

        for row in rows[1:]:
            item = make_item(
                source_file=source_file,
                index=index,
                question=row[question_idx] if question_idx < len(row) else "",
                answer=row[answer_idx] if answer_idx < len(row) else "",
                category=row[category_idx] if category_idx is not None and category_idx < len(row) else "",
                model=row[model_idx] if model_idx is not None and model_idx < len(row) else "",
                trim=row[trim_idx] if trim_idx is not None and trim_idx < len(row) else "",
                brand=row[brand_idx] if brand_idx is not None and brand_idx < len(row) else "",
                source_type="table",
            )
            if item:
                items.append(item)
                index += 1

        consumed_texts.extend(" ".join(row) for row in rows)

    return items, consumed_texts


def paragraph_texts(doc: Document) -> list[str]:
    return [
        clean_text(paragraph.text)
        for paragraph in doc.paragraphs
        if clean_text(paragraph.text)
    ]


def table_cell_texts(doc: Document) -> list[str]:
    values: list[str] = []
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = clean_text(cell.text)
                if text:
                    values.append(text)
    return values


def split_inline_qa(text: str) -> tuple[str, str] | None:
    patterns = [
        r"^(?:Q|q|问|问题)\s*[:：]\s*(?P<q>.+?)\s*(?:A|a|答|回答)\s*[:：]\s*(?P<a>.+)$",
        r"^(?P<q>.+?[?？])\s*(?:A|a|答|回答)\s*[:：]\s*(?P<a>.+)$",
    ]
    for pattern in patterns:
        match = re.match(pattern, text)
        if match:
            return match.group("q"), match.group("a")
    return None


def strip_number_prefix(text: str) -> str:
    return re.sub(r"^\s*(?:\d+[\.\、]|[（(]?\d+[）)]|[一二三四五六七八九十]+[、.])\s*", "", text)


def is_question_line(text: str) -> bool:
    stripped = strip_number_prefix(text)
    if re.match(r"^(?:Q|q|问|问题)\s*[:：]", stripped):
        return True
    return stripped.endswith(("?", "？"))


def extract_question(text: str) -> str:
    text = strip_number_prefix(text)
    return re.sub(r"^(?:Q|q|问|问题)\s*[:：]\s*", "", text).strip()


def is_answer_line(text: str) -> bool:
    return bool(re.match(r"^(?:A|a|答|回答)\s*[:：]", text))


def extract_answer(text: str) -> str:
    return re.sub(r"^(?:A|a|答|回答)\s*[:：]\s*", "", text).strip()


def restore_qa_paragraphs(lines: list[str], source_file: str) -> tuple[list[KnowledgeItem], set[int]]:
    items: list[KnowledgeItem] = []
    consumed: set[int] = set()
    index = 10000
    i = 0

    while i < len(lines):
        line = lines[i]
        inline = split_inline_qa(line)
        if inline:
            item = make_item(
                source_file=source_file,
                index=index,
                question=inline[0],
                answer=inline[1],
                source_type="inline_qa",
            )
            if item:
                items.append(item)
                consumed.add(i)
                index += 1
            i += 1
            continue

        if is_question_line(line):
            question = extract_question(line)
            answer_parts = []
            consumed_indexes = {i}
            j = i + 1
            while j < len(lines):
                next_line = lines[j]
                if is_question_line(next_line):
                    break
                if is_answer_line(next_line):
                    answer_parts.append(extract_answer(next_line))
                    consumed_indexes.add(j)
                elif answer_parts:
                    answer_parts.append(next_line)
                    consumed_indexes.add(j)
                else:
                    break
                j += 1

            if answer_parts:
                item = make_item(
                    source_file=source_file,
                    index=index,
                    question=question,
                    answer="\n".join(answer_parts),
                    source_type="qa_paragraph",
                )
                if item:
                    items.append(item)
                    consumed.update(consumed_indexes)
                    index += 1
                    i = j
                    continue
        i += 1

    return items, consumed


def likely_heading(text: str) -> bool:
    if looks_like_json_text(text) or contains_json_marker(text):
        return False
    if len(text) > 40:
        return False
    if text.endswith(("。", "，", ",", "；", ";")):
        return False
    if is_answer_line(text):
        return False
    return bool(re.search(r"(政策|权益|价格|金融|配置|质保|救援|流量|续航|动力|空间|智驾|车型|基础信息)", text))


def restore_heading_blocks(lines: list[str], consumed: set[int], source_file: str) -> tuple[list[KnowledgeItem], set[int]]:
    items: list[KnowledgeItem] = []
    heading_consumed: set[int] = set()
    index = 20000
    i = 0

    while i < len(lines):
        if i in consumed or not likely_heading(lines[i]):
            i += 1
            continue

        heading = lines[i]
        body_parts = []
        used = {i}
        j = i + 1
        while j < len(lines):
            if j in consumed:
                j += 1
                continue
            if likely_heading(lines[j]) or is_question_line(lines[j]):
                break
            body_parts.append(lines[j])
            used.add(j)
            j += 1

        body = "\n".join(body_parts).strip()
        if body and len(body) >= 8:
            item = make_item(
                source_file=source_file,
                index=index,
                question=heading,
                answer=body,
                source_type="heading_body",
            )
            if item:
                items.append(item)
                heading_consumed.update(used)
                index += 1
                i = j
                continue
        i += 1

    return items, heading_consumed


def dedupe_items(items: list[KnowledgeItem]) -> tuple[list[KnowledgeItem], list[KnowledgeItem]]:
    kept: list[KnowledgeItem] = []
    uncertain: list[KnowledgeItem] = []
    seen: dict[tuple[str, str, str, str], KnowledgeItem] = {}

    for item in items:
        key = (
            item.model or "",
            item.trim or "",
            item.category or "",
            item.normalized_question,
        )
        existing = seen.get(key)
        if not existing:
            seen[key] = item
            kept.append(item)
            continue

        if clean_text(existing.answer) == clean_text(item.answer):
            continue

        item.need_confirm = True
        item.metadata["duplicate_conflict_with"] = existing.knowledge_id
        uncertain.append(item)

    return kept + uncertain, uncertain


def restore_word_fast(file: Any) -> FastRestoreResult:
    source_file = get_source_name(file)
    doc = load_document(file)
    lines = paragraph_texts(doc)
    cell_lines = table_cell_texts(doc)

    table_items, table_consumed_texts = restore_tables(doc, source_file)
    json_items, json_consumed, json_failed = restore_json_blocks(lines + cell_lines, source_file)
    qa_items, qa_consumed = restore_qa_paragraphs(lines, source_file)
    qa_consumed = set(qa_consumed) - set(json_consumed)
    heading_items, heading_consumed = restore_heading_blocks(lines, set(qa_consumed) | set(json_consumed), source_file)

    consumed = set(json_consumed) | set(qa_consumed) | set(heading_consumed)
    remaining_lines = [
        line
        for index, line in enumerate(lines)
        if index not in consumed
    ]
    table_text = "\n".join(table_consumed_texts)
    remaining_text = "\n".join(remaining_lines).strip()
    document_text = "\n".join(lines + ([table_text] if table_text else []))
    items, uncertain = dedupe_items(json_items + table_items + qa_items + heading_items)
    attempted = len(items) + json_failed + len(uncertain)
    confidence = (len(items) / attempted) if attempted else 0.0

    return FastRestoreResult(
        items=items,
        remaining_text=remaining_text,
        document_text=document_text,
        total_chars=len(document_text),
        remaining_chars=len(remaining_text),
        consumed_blocks=len(consumed) + len(table_consumed_texts),
        table_items=len(table_items),
        qa_items=len(qa_items),
        heading_items=len(heading_items),
        json_items=len(json_items),
        failed_items=json_failed + len(uncertain),
        confidence=confidence,
    )
