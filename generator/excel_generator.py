import os
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from utils.rag_quality import (
    apply_export_quality_gate,
    is_exportable_rag,
    normalize_answer_for_tts
)


DYNAMIC_KEYWORDS = [
    "价格",
    "售价",
    "金融",
    "优惠",
    "权益",
    "活动",
    "政策",
    "补贴"
]


def format_sheet(ws):

    ws.freeze_panes = "A2"

    if ws.max_row > 1:

        ws.auto_filter.ref = ws.dimensions

    for cell in ws[1]:

        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    for row in ws.iter_rows():

        for cell in row:

            cell.alignment = Alignment(
                wrap_text=True,
                vertical="top"
            )

    width_config = {

        "A": 18,

        "B": 18,

        "C": 36,

        "D": 70,

        "E": 18

    }

    for col, width in width_config.items():

        ws.column_dimensions[
            col
        ].width = width


def infer_knowledge_type(item):

    knowledge_type = item.get(
        "knowledge_type",
        ""
    )

    if knowledge_type in [
        "static",
        "dynamic"
    ]:

        return knowledge_type

    text = (
        f"{item.get('category', '')} "
        f"{item.get('module', '')} "
        f"{item.get('answer', '')}"
    )

    if any(
        keyword in text
        for keyword in DYNAMIC_KEYWORDS
    ):

        return "dynamic"

    return "static"


def get_questions(item):

    questions = item.get(
        "questions",
        item.get(
            "question",
            []
        )
    )

    if isinstance(
        questions,
        list
    ):

        return "\n".join(
            str(question)
            for question in questions
            if question
        )

    return str(
        questions or ""
    )


def get_category(item):

    return (
        item.get(
            "category",
            ""
        )
        or
        item.get(
            "module",
            ""
        )
    )


def append_rows(ws, items):

    ws.append(
        [
            "车型",
            "版本",
            "问题",
            "回答",
            "分类"
        ]
    )

    for item in items:

        ws.append(
            [
                item.get(
                    "model",
                    ""
                ),
                item.get(
                    "trim",
                    ""
                ),
                get_questions(
                    item
                ),
                item.get(
                    "answer",
                    ""
                ),
                get_category(
                    item
                )
            ]
        )

    format_sheet(
        ws
    )


def save_simple_workbook(items, sheet_name, output_path):

    output_dir = os.path.dirname(
        output_path
    )

    if output_dir:

        os.makedirs(
            output_dir,
            exist_ok=True
        )

    wb = Workbook()

    ws = wb.active

    ws.title = sheet_name

    append_rows(
        ws,
        items
    )

    wb.save(
        output_path
    )


def clean_filename(value):

    value = str(
        value or ""
    ).strip()

    value = re.sub(
        r'[\\/:*?"<>|]',
        "",
        value
    )

    return value


def normalize_brand_model_name(brand, model):

    brand = clean_filename(
        brand
    )

    model = clean_filename(
        model
    )

    if brand and model.startswith(
        brand
    ):

        return model

    return clean_filename(
        f"{brand}{model}"
    )


def infer_project_name(facts_data, rag_data, output_path):

    candidates = []

    candidates.extend(
        rag_data.get(
            "rag_knowledge",
            []
        )
    )

    candidates.extend(
        facts_data.get(
            "facts",
            []
        )
    )

    brands = []

    models = []

    for item in candidates:

        if not isinstance(
            item,
            dict
        ):

            continue

        brand = item.get(
            "brand",
            ""
        )

        model = item.get(
            "model",
            ""
        )

        if brand and brand not in brands:

            brands.append(
                brand
            )

        if model and model not in models:

            models.append(
                model
            )

    if len(
        brands
    ) == 1 and len(
        models
    ) == 1:

        return clean_filename(
            normalize_brand_model_name(
                brands[0],
                models[0]
            )
        )

    if len(
        brands
    ) == 1 and len(
        models
    ) > 1:

        return clean_filename(
            brands[0]
        )

    if len(
        brands
    ) > 1:

        return "多品牌"

    base_name = os.path.splitext(
        os.path.basename(
            output_path
        )
    )[0]

    return clean_filename(
        base_name
    )


def build_output_paths(output_path, facts_data, rag_data):

    output_dir = os.path.dirname(
        output_path
    )

    base_name = infer_project_name(
        facts_data,
        rag_data,
        output_path
    )

    static_path = os.path.join(
        output_dir,
        f"{base_name}_车型配置知识库.xlsx"
    )

    dynamic_path = os.path.join(
        output_dir,
        f"{base_name}_价格政策知识库.xlsx"
    )

    return {
        "static": static_path,
        "dynamic": dynamic_path
    }


def generate_excel(
        facts_data,
        rag_data,
        qc_data,
        output_path
):

    rag_list = rag_data.get(
        "rag_knowledge",
        []
    )

    apply_export_quality_gate(
        rag_data
    )

    static_items = []

    dynamic_items = []

    excluded_items = []

    for item in rag_list:

        if not isinstance(
            item,
            dict
        ):

            continue

        item[
            "answer"
        ] = normalize_answer_for_tts(
            item.get(
                "answer",
                ""
            )
        )

        if not is_exportable_rag(
            item
        ):

            excluded_items.append(
                item
            )

            continue

        if infer_knowledge_type(
            item
        ) == "dynamic":

            dynamic_items.append(
                item
            )

        else:

            static_items.append(
                item
            )

    rag_data[
        "export_excluded"
    ] = excluded_items

    rag_data[
        "export_excluded_count"
    ] = len(
        excluded_items
    )

    output_paths = build_output_paths(
        output_path,
        facts_data,
        rag_data
    )

    save_simple_workbook(
        static_items,
        "车型配置知识库",
        output_paths[
            "static"
        ]
    )

    save_simple_workbook(
        dynamic_items,
        "价格政策知识库",
        output_paths[
            "dynamic"
        ]
    )

    return output_paths
