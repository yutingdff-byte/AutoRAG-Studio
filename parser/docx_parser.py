import os
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

from docx import Document
from PIL import Image, UnidentifiedImageError

from parser.image_parser import parse_image


IMAGE_DOMINANT_TEXT_CHAR_THRESHOLD = 200
MIN_EMBEDDED_IMAGE_WIDTH = 240
MIN_EMBEDDED_IMAGE_HEIGHT = 120
MIN_EMBEDDED_IMAGE_AREA = 50_000


@dataclass
class EmbeddedWordImage:
    name: str
    data: bytes
    width: int
    height: int
    index: int

    def getvalue(self) -> bytes:
        return self.data


def _get_file_name(file) -> str:
    if isinstance(file, str):
        return os.path.basename(file)
    return Path(str(getattr(file, "name", "uploaded.docx"))).name


def _load_document(file):
    if not isinstance(file, str) and hasattr(file, "seek"):
        try:
            file.seek(0)
        except Exception:
            pass
    return Document(file)


def _extract_text_blocks(doc, file_name: str) -> list[str]:
    texts = [f"【文件来源：{file_name}】"]

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            texts.append(text)

    for index, table in enumerate(doc.tables):
        table_rows = []
        for row in table.rows:
            row_data = [cell.text.strip() for cell in row.cells]
            if any(row_data):
                table_rows.append(" | ".join(row_data))

        if table_rows:
            texts.append(f"\n【表格 {index + 1}】")
            texts.extend(table_rows)

    return texts


def _effective_text_length(text_blocks: list[str]) -> int:
    content = "\n".join(text_blocks[1:]) if len(text_blocks) > 1 else ""
    return len("".join(content.split()))


def _image_suffix(content_type: str) -> str:
    mapping = {
        "image/png": ".png",
        "image/jpeg": ".jpg",
        "image/jpg": ".jpg",
        "image/webp": ".webp",
    }
    return mapping.get(content_type.lower(), ".png")


def _is_valid_content_image(width: int, height: int) -> bool:
    return (
        width >= MIN_EMBEDDED_IMAGE_WIDTH
        and height >= MIN_EMBEDDED_IMAGE_HEIGHT
        and width * height >= MIN_EMBEDDED_IMAGE_AREA
    )


def _extract_embedded_images(doc, file_name: str) -> list[EmbeddedWordImage]:
    images = []
    seen_rids = set()

    for element in doc.element.body.iter():
        for blip in element.xpath(".//*[local-name()='blip']"):
            rel_id = blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
            if not rel_id or rel_id in seen_rids:
                continue
            seen_rids.add(rel_id)

            part = doc.part.related_parts.get(rel_id)
            if not part:
                continue

            data = part.blob
            try:
                with Image.open(BytesIO(data)) as image:
                    width, height = image.size
            except (UnidentifiedImageError, OSError):
                continue

            if not _is_valid_content_image(width, height):
                continue

            index = len(images) + 1
            suffix = _image_suffix(getattr(part, "content_type", "image/png"))
            images.append(
                EmbeddedWordImage(
                    name=f"{Path(file_name).stem}_image_{index:03d}{suffix}",
                    data=data,
                    width=width,
                    height=height,
                    index=index,
                )
            )

    return images


def _render_text_only(text_blocks: list[str]) -> str:
    return "\n".join(text_blocks)


def _render_image_dominant_word(file_name: str, text_blocks: list[str], images: list[EmbeddedWordImage]) -> str:
    result = [
        text_blocks[0],
        "检测到该 Word 主要由图片组成，正在使用图片识别。",
    ]

    body_text = "\n".join(text_blocks[1:]).strip()
    if body_text:
        result.extend(["\n【Word正文】", body_text])

    success_count = 0
    failed_messages = []
    total_images = len(images)

    for image in images:
        result.append(f"\n【内嵌图片 {image.index}/{total_images}：{image.name}】")
        try:
            recognized_text = parse_image(image).strip()
            if recognized_text:
                result.append(recognized_text)
                success_count += 1
            else:
                message = f"图片 {image.index}/{total_images} 识别失败：未返回有效内容"
                result.append(f"【{message}】")
                failed_messages.append(message)
        except Exception as exc:
            message = f"图片 {image.index}/{total_images} 识别失败：{exc}"
            result.append(f"【{message}】")
            failed_messages.append(message)

    if success_count == 0:
        raise ValueError("未能从该 Word 中识别到有效内容，请检查文件内容或重新上传。")

    if failed_messages:
        result.append(f"\n【图片识别提示】成功 {success_count}/{total_images}，失败 {len(failed_messages)}/{total_images}。")

    return "\n".join(result)


def parse_docx(file):
    """
    Word文档解析器。

    文本型 Word 保持原文本解析行为；当文档正文很少且包含有效大图时，
    自动复用现有图片解析能力提取内嵌图片内容。
    """

    try:
        file_name = _get_file_name(file)
        doc = _load_document(file)
        text_blocks = _extract_text_blocks(doc, file_name)
        valid_images = _extract_embedded_images(doc, file_name)

        is_image_dominant = (
            _effective_text_length(text_blocks) < IMAGE_DOMINANT_TEXT_CHAR_THRESHOLD
            and bool(valid_images)
        )

        if is_image_dominant:
            return _render_image_dominant_word(file_name, text_blocks, valid_images)

        return _render_text_only(text_blocks)

    except Exception as e:
        raise Exception(f"Word解析失败: {str(e)}")


if __name__ == "__main__":
    result = parse_docx("test.docx")
    print(result)
