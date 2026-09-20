from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any, BinaryIO

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.worksheet.worksheet import Worksheet
import xlrd


SUPPORTED_EXTENSIONS = {
    ".xlsx",
    ".xlsm",
    ".xltx",
    ".xltm",
}

XLRD_EXTENSIONS = {
    ".xls",
}

ROW_ORIENTED_TABLE = "ROW_ORIENTED_TABLE"
COLUMN_ORIENTED_VEHICLE_MATRIX = "COLUMN_ORIENTED_VEHICLE_MATRIX"
UNKNOWN_EXCEL_ORIENTATION = "UNKNOWN"

MATRIX_IDENTITY_FIELDS = {
    "品牌",
    "车系名称",
    "车型名称",
    "配置版本等级",
    "厂商指导价",
}

MATRIX_ATTRIBUTE_SIGNALS = MATRIX_IDENTITY_FIELDS | {
    "车型",
    "一句话介绍",
    "车型价格区间",
    "细分品牌分类",
    "车辆级别",
    "能源类型",
    "购车补贴政策",
    "厂家置换补贴",
    "金融优惠",
    "其他权益",
    "其他活动",
}

MIN_MATRIX_IDENTITY_ROWS = 3
MIN_MATRIX_VEHICLE_COLUMNS = 2
MIN_MATRIX_DATA_VALUES_PER_COLUMN = 3


def _clean_text(value: Any) -> str:
    """
    将Excel单元格值转换成适合传给大模型的文本。
    """

    if value is None:
        return ""

    if isinstance(value, bool):
        return "是" if value else "否"

    if isinstance(value, datetime):
        if value.time().hour == 0 and value.time().minute == 0:
            return value.strftime("%Y-%m-%d")
        return value.strftime("%Y-%m-%d %H:%M:%S")

    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")

    if isinstance(value, float):
        # 249800.0 -> 249800
        if value.is_integer():
            return str(int(value))

        # 避免出现过长的小数
        return str(round(value, 8)).rstrip("0").rstrip(".")

    text = str(value).strip()

    # 清理常见不可见字符
    text = text.replace("\u00a0", " ")
    text = text.replace("\u3000", " ")
    text = text.replace("\r\n", "\n")
    text = text.replace("\r", "\n")

    return text.strip()


def _get_source_name(file_path: Any) -> str:
    """
    获取文件展示名称。
    兼容路径和类似Streamlit UploadedFile的对象。
    """

    name = getattr(file_path, "name", None)

    if name:
        return Path(str(name)).name

    if isinstance(file_path, (str, Path)):
        return Path(file_path).name

    return "uploaded_excel.xlsx"


def _prepare_source(file_path: Any) -> Any:
    """
    将输入转换为openpyxl可以读取的对象。
    """

    if isinstance(file_path, bytes):
        return BytesIO(file_path)

    if isinstance(file_path, bytearray):
        return BytesIO(bytes(file_path))

    if hasattr(file_path, "read"):
        try:
            file_path.seek(0)
        except Exception:
            pass

        return file_path

    return file_path


def _validate_extension(file_path: Any) -> None:
    """
    当能够识别扩展名时，检查是否为支持的Excel格式。
    """

    name = getattr(file_path, "name", None)

    if name is None and isinstance(file_path, (str, Path)):
        name = str(file_path)

    if not name:
        return

    suffix = Path(str(name)).suffix.lower()

    if suffix and suffix not in SUPPORTED_EXTENSIONS and suffix not in XLRD_EXTENSIONS:
        raise ValueError(
            f"暂不支持该Excel格式：{suffix}。"
            "请将文件另存为 .xlsx 或上传标准 .xls 文件后重试。"
        )


def _get_suffix(file_path: Any) -> str:
    name = getattr(file_path, "name", None)

    if name is None and isinstance(file_path, (str, Path)):
        name = str(file_path)

    if not name:
        return ""

    return Path(str(name)).suffix.lower()


def _build_merged_value_map(ws: Worksheet) -> dict[str, Any]:
    """
    将合并单元格区域内的所有坐标映射到左上角单元格的值。

    例如：
    A1:C1合并，只有A1有值。
    解析时B1、C1也能识别为A1的值。
    """

    merged_value_map: dict[str, Any] = {}

    for merged_range in ws.merged_cells.ranges:
        min_col, min_row, max_col, max_row = merged_range.bounds

        top_left_value = ws.cell(
            row=min_row,
            column=min_col,
        ).value

        for row_index in range(min_row, max_row + 1):
            for col_index in range(min_col, max_col + 1):
                coordinate = ws.cell(
                    row=row_index,
                    column=col_index,
                ).coordinate

                merged_value_map[coordinate] = top_left_value

    return merged_value_map


def _read_sheet_rows(ws: Worksheet) -> list[list[str]]:
    """
    读取Sheet中的有效行。
    """

    merged_value_map = _build_merged_value_map(ws)

    raw_rows: list[list[str]] = []

    for row in ws.iter_rows():
        values: list[str] = []

        for cell in row:
            if isinstance(cell, MergedCell):
                raw_value = merged_value_map.get(
                    cell.coordinate,
                    "",
                )
            else:
                raw_value = merged_value_map.get(
                    cell.coordinate,
                    cell.value,
                )

            values.append(
                _clean_text(raw_value)
            )

        # 删除行尾空列
        while values and not values[-1]:
            values.pop()

        # 整行为空则跳过
        if not any(values):
            continue

        raw_rows.append(values)

    return raw_rows


def _clean_xlrd_cell(book: xlrd.book.Book, sheet: xlrd.sheet.Sheet, row_index: int, col_index: int) -> str:
    cell = sheet.cell(row_index, col_index)
    value = cell.value

    if cell.ctype == xlrd.XL_CELL_EMPTY:
        return ""

    if cell.ctype == xlrd.XL_CELL_DATE:
        try:
            return _clean_text(
                xlrd.xldate.xldate_as_datetime(
                    value,
                    book.datemode,
                )
            )
        except Exception:
            return _clean_text(value)

    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return "是" if bool(value) else "否"

    if cell.ctype == xlrd.XL_CELL_NUMBER:
        return _clean_text(float(value))

    return _clean_text(value)


def _read_xls_sheet_rows(book: xlrd.book.Book, sheet: xlrd.sheet.Sheet) -> list[list[str]]:
    raw_rows: list[list[str]] = []

    for row_index in range(sheet.nrows):
        values: list[str] = []
        for col_index in range(sheet.ncols):
            values.append(
                _clean_xlrd_cell(
                    book,
                    sheet,
                    row_index,
                    col_index,
                )
            )

        while values and not values[-1]:
            values.pop()

        if not any(values):
            continue

        raw_rows.append(values)

    return raw_rows


def _count_non_empty(row: list[str]) -> int:
    return sum(
        1
        for value in row
        if value
    )


def _normalize_field_name(value: str) -> str:
    return (
        value
        .replace(" ", "")
        .replace("\n", "")
        .replace("\t", "")
        .strip()
    )


def _row_value(row: list[str], index: int) -> str:
    if index < len(row):
        return row[index]

    return ""


def _first_column_fields(rows: list[list[str]]) -> list[str]:
    return [
        _normalize_field_name(row[0])
        for row in rows
        if row and row[0]
    ]


def _identity_row_indexes(rows: list[list[str]]) -> dict[str, int]:
    indexes: dict[str, int] = {}

    for index, row in enumerate(rows):
        if not row:
            continue

        field_name = _normalize_field_name(row[0])

        if field_name in MATRIX_IDENTITY_FIELDS:
            indexes[field_name] = index

    return indexes


def _effective_column_count(rows: list[list[str]]) -> int:
    count = 0

    for row in rows:
        for index, value in enumerate(row):
            if value:
                count = max(count, index + 1)

    return count


def _matrix_vehicle_columns(
    rows: list[list[str]],
    identity_rows: dict[str, int],
) -> list[int]:
    """
    找出真正包含车型/版本数据的列。

    返回值是0-based列索引。第0列是属性名列，不会返回。
    """

    max_columns = _effective_column_count(rows)
    vehicle_columns: list[int] = []

    for column_index in range(1, max_columns):
        identity_values = [
            _row_value(rows[row_index], column_index)
            for row_index in identity_rows.values()
            if row_index < len(rows)
        ]

        identity_non_empty = _count_non_empty(identity_values)

        if identity_non_empty < 2:
            continue

        data_non_empty = 0

        for row in rows:
            if _row_value(row, column_index):
                data_non_empty += 1

        if data_non_empty < MIN_MATRIX_DATA_VALUES_PER_COLUMN:
            continue

        vehicle_columns.append(column_index)

    return vehicle_columns


def _detect_excel_orientation(rows: list[list[str]]) -> str:
    """
    识别Excel表格方向。

    ROW_ORIENTED_TABLE保留原有行记录表逻辑；
    COLUMN_ORIENTED_VEHICLE_MATRIX表示A列是属性、B之后每列是车型版本。
    """

    if not rows:
        return UNKNOWN_EXCEL_ORIENTATION

    first_column = _first_column_fields(rows)

    if len(first_column) < 5:
        return ROW_ORIENTED_TABLE

    signal_hits = sum(
        1
        for field_name in first_column
        if field_name in MATRIX_ATTRIBUTE_SIGNALS
    )

    identity_rows = _identity_row_indexes(rows)

    vehicle_columns = _matrix_vehicle_columns(
        rows,
        identity_rows,
    )

    first_column_signal_ratio = signal_hits / len(first_column)

    if (
        len(identity_rows) >= MIN_MATRIX_IDENTITY_ROWS
        and len(vehicle_columns) >= MIN_MATRIX_VEHICLE_COLUMNS
        and first_column_signal_ratio >= 0.25
    ):
        return COLUMN_ORIENTED_VEHICLE_MATRIX

    return ROW_ORIENTED_TABLE


def _looks_like_title_row(row: list[str]) -> bool:
    """
    判断一行是否更像文件标题，而不是表头。
    """

    non_empty_values = [
        value
        for value in row
        if value
    ]

    if len(non_empty_values) != 1:
        return False

    value = non_empty_values[0]

    # 单个长文本通常是标题或说明
    return len(value) >= 6


def _detect_header_index(rows: list[list[str]]) -> int | None:
    """
    在前20个有效行中寻找最可能的表头。

    典型格式：

    第1行：阿维塔07用户购车权益
    第2行：活动时间说明
    第3行：序号 | 权益名称 | 权益内容 | 适用版本

    返回第3行对应的索引。
    """

    if not rows:
        return None

    search_limit = min(
        len(rows),
        20,
    )

    best_index: int | None = None
    best_score = float("-inf")

    header_keywords = {
        "车型",
        "品牌",
        "版本",
        "配置",
        "参数",
        "价格",
        "指导价",
        "权益",
        "政策",
        "内容",
        "名称",
        "序号",
        "项目",
        "说明",
        "时间",
        "日期",
        "适用范围",
        "适用版本",
        "备注",
        "续航",
        "轴距",
    }

    for index in range(search_limit):
        row = rows[index]

        non_empty = [
            value
            for value in row
            if value
        ]

        non_empty_count = len(non_empty)

        if non_empty_count < 2:
            continue

        score = non_empty_count * 3

        # 表头通常以文本字段名为主
        text_count = sum(
            1
            for value in non_empty
            if not value.replace(".", "", 1).isdigit()
        )
        score += text_count

        # 命中常见字段名称
        keyword_hits = 0

        for value in non_empty:
            compact_value = value.replace(" ", "")

            if any(
                keyword in compact_value
                for keyword in header_keywords
            ):
                keyword_hits += 1

        score += keyword_hits * 4

        # 下一行存在较多数据，当前行更可能是表头
        if index + 1 < len(rows):
            next_non_empty = _count_non_empty(
                rows[index + 1]
            )

            if next_non_empty >= 2:
                score += min(
                    next_non_empty,
                    non_empty_count,
                ) * 2

        # 单元格内容过长，可能是正文或说明
        long_cell_count = sum(
            1
            for value in non_empty
            if len(value) > 40
        )
        score -= long_cell_count * 3

        # 单行大段文字不像表头
        if _looks_like_title_row(row):
            score -= 10

        if score > best_score:
            best_score = score
            best_index = index

    return best_index


def _make_unique_headers(
    header_row: list[str],
    max_columns: int,
) -> list[str]:
    """
    补全空表头，并处理重名字段。
    """

    headers: list[str] = []
    name_counts: dict[str, int] = {}

    for index in range(max_columns):
        if index < len(header_row):
            header = header_row[index].strip()
        else:
            header = ""

        if not header:
            header = f"字段{index + 1}"

        count = name_counts.get(
            header,
            0,
        ) + 1

        name_counts[header] = count

        if count > 1:
            unique_header = f"{header}_{count}"
        else:
            unique_header = header

        headers.append(unique_header)

    return headers


def _format_preface_rows(
    rows: list[list[str]],
) -> list[str]:
    """
    格式化表头之前的标题、活动说明等信息。
    """

    output: list[str] = []

    for row in rows:
        values = [
            value
            for value in row
            if value
        ]

        if not values:
            continue

        output.append(
            " | ".join(values)
        )

    return output


def _format_structured_sheet(
    rows: list[list[str]],
    header_index: int,
) -> list[str]:
    """
    将带表头的数据格式化为：

    字段：车型 | 续航 | 轴距 | 价格

    记录1：
    车型: 理想i6
    续航: 720km
    ...
    """

    output: list[str] = []

    preface_rows = rows[:header_index]

    if preface_rows:
        preface_text = _format_preface_rows(
            preface_rows
        )

        if preface_text:
            output.append("【前置信息】")
            output.extend(preface_text)
            output.append("")

    data_rows = rows[header_index + 1:]

    max_columns = max(
        [len(rows[header_index])]
        + [
            len(row)
            for row in data_rows
        ]
    )

    headers = _make_unique_headers(
        rows[header_index],
        max_columns,
    )

    output.append(
        "字段："
        + " | ".join(headers)
    )

    record_number = 0

    for row in data_rows:
        fields: list[str] = []

        for index, header in enumerate(headers):
            value = (
                row[index]
                if index < len(row)
                else ""
            )

            if not value:
                continue

            fields.append(
                f"{header}: {value}"
            )

        if not fields:
            continue

        record_number += 1

        output.append("")
        output.append(
            f"记录{record_number}："
        )
        output.extend(fields)

    return output


def _format_column_oriented_vehicle_matrix(
    rows: list[list[str]],
) -> list[str]:
    """
    将横向车型矩阵确定性转置为一列一个车型版本记录。

    原始结构：
        A列 = 属性名
        B之后 = 车型/版本

    输出结构：
        【车型记录1】
        字段: 值
    """

    output: list[str] = [
        "【表格结构：横向车型矩阵】"
    ]

    identity_rows = _identity_row_indexes(rows)
    vehicle_columns = _matrix_vehicle_columns(
        rows,
        identity_rows,
    )

    record_number = 0

    for column_index in vehicle_columns:
        fields: list[str] = []
        seen_fields: dict[str, int] = {}

        for row in rows:
            if not row:
                continue

            raw_field_name = _row_value(
                row,
                0,
            )
            value = _row_value(
                row,
                column_index,
            )

            if not raw_field_name or not value:
                continue

            field_name = raw_field_name.strip()
            count = seen_fields.get(
                field_name,
                0,
            ) + 1
            seen_fields[field_name] = count

            if count > 1:
                field_name = f"{field_name}_{count}"

            fields.append(
                f"{field_name}: {value}"
            )

        if not fields:
            continue

        record_number += 1
        output.append("")
        output.append(
            f"【车型记录{record_number}】"
        )
        output.extend(fields)

    return output


def _format_unstructured_sheet(
    rows: list[list[str]],
) -> list[str]:
    """
    没有明显表头时，按列序号输出，避免丢失内容。
    """

    output: list[str] = []

    for record_number, row in enumerate(
        rows,
        start=1,
    ):
        fields: list[str] = []

        for column_index, value in enumerate(
            row,
            start=1,
        ):
            if not value:
                continue

            fields.append(
                f"列{column_index}: {value}"
            )

        if not fields:
            continue

        output.append(
            f"记录{record_number}："
        )
        output.extend(fields)
        output.append("")

    return output


def _format_sheet_rows(rows: list[list[str]]) -> list[str]:
    orientation = _detect_excel_orientation(
        rows
    )

    if orientation == COLUMN_ORIENTED_VEHICLE_MATRIX:
        return _format_column_oriented_vehicle_matrix(
            rows
        )

    header_index = _detect_header_index(
        rows
    )

    if header_index is not None:
        return _format_structured_sheet(
            rows,
            header_index,
        )

    return _format_unstructured_sheet(
        rows
    )


def parse_excel(
    file_path: str | Path | BinaryIO | bytes,
) -> str:
    """
    解析Excel文件并返回统一文本。

    参数：
        file_path:
            文件路径、二进制文件对象、UploadedFile或bytes。

    返回：
        可直接合并到统一material中的文本。
    """

    _validate_extension(file_path)

    source_name = _get_source_name(
        file_path
    )

    suffix = _get_suffix(
        file_path
    )

    source = _prepare_source(
        file_path
    )

    if suffix in XLRD_EXTENSIONS:
        try:
            if isinstance(source, (str, Path)):
                workbook = xlrd.open_workbook(
                    filename=str(source),
                )
            else:
                if hasattr(source, "seek"):
                    source.seek(0)
                data = source.read() if hasattr(source, "read") else source
                workbook = xlrd.open_workbook(
                    file_contents=bytes(data),
                )
        except Exception as exc:
            raise RuntimeError(
                f"Excel解析失败：{source_name}；"
                f"错误信息：{exc}"
            ) from exc

        output: list[str] = [
            f"【文件来源：{source_name}】"
        ]

        parsed_sheet_count = 0

        for sheet in workbook.sheets():
            rows = _read_xls_sheet_rows(
                workbook,
                sheet,
            )

            if not rows:
                continue

            parsed_sheet_count += 1

            output.append("")
            output.append(
                f"【Sheet:{sheet.name}】"
            )
            output.extend(
                _format_sheet_rows(
                    rows
                )
            )

        if parsed_sheet_count == 0:
            output.append("")
            output.append(
                "【提示】该Excel中未读取到有效内容。"
            )

        return "\n".join(output).strip()

    try:
        workbook = load_workbook(
            filename=source,
            data_only=True,
            read_only=False,
        )

    except Exception as exc:
        raise RuntimeError(
            f"Excel解析失败：{source_name}；"
            f"错误信息：{exc}"
        ) from exc

    output: list[str] = [
        f"【文件来源：{source_name}】"
    ]

    parsed_sheet_count = 0

    for ws in workbook.worksheets:
        rows = _read_sheet_rows(ws)

        # 跳过完全空白的Sheet
        if not rows:
            continue

        parsed_sheet_count += 1

        output.append("")
        output.append(
            f"【Sheet:{ws.title}】"
        )

        output.extend(
            _format_sheet_rows(
                rows
            )
        )

    workbook.close()

    if parsed_sheet_count == 0:
        output.append("")
        output.append(
            "【提示】该Excel中未读取到有效内容。"
        )

    return "\n".join(output).strip()


if __name__ == "__main__":
    # 本地测试：
    # python parser/excel_parser.py
    test_file = "tests/data/test.xlsx"

    try:
        result = parse_excel(test_file)
        print(result)

    except Exception as error:
        print("Excel测试解析失败：")
        print(error)
