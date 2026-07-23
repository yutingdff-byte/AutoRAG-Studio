import os
import tempfile
from pathlib import Path

from openpyxl import load_workbook

from generator.excel_generator import generate_excel
from parser.image_parser import IMAGE_EXTENSIONS
from utils.config import get_config


def test_config_reads_environment(monkeypatch):

    monkeypatch.setenv(
        "AUTORAG_TEST_VALUE",
        "ok"
    )

    assert get_config(
        "AUTORAG_TEST_VALUE"
    ) == "ok"


def test_excel_quality_gate_keeps_columns_and_blocks_missing():

    facts_data = {
        "facts": [
            {
                "brand": "阿维塔",
                "model": "阿维塔07"
            }
        ]
    }

    rag_data = {
        "rag_knowledge": [
            {
                "rag_id": "RAG-001",
                "brand": "阿维塔",
                "model": "阿维塔07",
                "trim": "全系",
                "category": "空间",
                "knowledge_type": "static",
                "answer_type": "sales_translation",
                "questions": [
                    "空间怎么样？"
                ],
                "answer": "空间比较宽裕，家庭日常使用会比较舒服。",
                "fact_refs": [
                    "F001"
                ]
            },
            {
                "rag_id": "RAG-002",
                "brand": "阿维塔",
                "model": "阿维塔07",
                "trim": "全系",
                "category": "价格",
                "knowledge_type": "dynamic",
                "answer_type": "fact_answer",
                "questions": [
                    "多少钱？"
                ],
                "answer": "官方指导价21.99万元起，具体成交价格以门店为准。",
                "fact_refs": [
                    "F002"
                ]
            },
            {
                "rag_id": "RAG-003",
                "brand": "阿维塔",
                "model": "阿维塔07",
                "trim": "全系",
                "category": "金融",
                "knowledge_type": "static",
                "answer_type": "need_confirm",
                "need_confirm": "是",
                "questions": [
                    "有没有金融方案？"
                ],
                "answer": "目前没有相关信息，建议咨询门店。",
                "fact_refs": []
            }
        ]
    }

    with tempfile.TemporaryDirectory() as tmpdir:

        paths = generate_excel(
            facts_data,
            rag_data,
            {},
            str(
                Path(
                    tmpdir
                )
                / "result.xlsx"
            )
        )

        row_counts = {}

        for key, path in paths.items():

            workbook = load_workbook(
                path,
                read_only=True,
                data_only=True
            )

            sheet = workbook.active

            headers = [
                cell.value
                for cell in sheet[1]
            ]

            rows = list(
                sheet.iter_rows(
                    min_row=2,
                    values_only=True
                )
            )

            row_counts[
                key
            ] = len(
                rows
            )

            assert headers == [
                "车型",
                "版本",
                "问题",
                "回答",
                "分类"
            ]

            assert all(
                "目前没有" not in str(
                    row[3]
                )
                for row in rows
            )

            workbook.close()

        assert row_counts == {
            "static": 1,
            "dynamic": 1
        }

        assert rag_data[
            "export_excluded_count"
        ] == 1


def test_image_extension_set_is_cloud_safe():

    assert {
        ".png",
        ".jpg",
        ".jpeg",
        ".webp"
    }.issubset(
        IMAGE_EXTENSIONS
    )


def test_temp_paths_are_unique():

    first = tempfile.TemporaryDirectory()
    second = tempfile.TemporaryDirectory()

    try:

        assert first.name != second.name

        assert os.path.isabs(
            first.name
        )

    finally:

        first.cleanup()
        second.cleanup()
