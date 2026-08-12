from io import BytesIO

from openpyxl import load_workbook

from knowledge.deduplicator import cleanup_exact_duplicates
from knowledge.models import KnowledgeItem
from pages.update_page import _build_update_export_files


def tank500_range_item(knowledge_id: str) -> KnowledgeItem:
    return KnowledgeItem(
        knowledge_id=knowledge_id,
        model="坦克500",
        trim="",
        question="坦克500的续航",
        answer="坦克500的WLTC纯电续航有110公里和201公里两个版本的，综合算下来的话有900公里和1096公里两个版本",
        category="",
        knowledge_type="product",
    )


def read_export_rows(data: bytes) -> list[dict]:
    workbook = load_workbook(BytesIO(data))
    worksheet = workbook.active
    headers = [cell.value for cell in worksheet[1]]
    return [
        dict(zip(headers, row))
        for row in worksheet.iter_rows(min_row=2, values_only=True)
    ]


def test_update_export_uses_cleaned_final_knowledge_only():
    cleanup = cleanup_exact_duplicates(
        [
            tank500_range_item("STD-DOCX-f442891cc2a3"),
            tank500_range_item("STD-DOCX-d6d9017dd16b"),
        ]
    )

    files = _build_update_export_files(cleanup.final_items)
    static_rows = read_export_rows(files["static"]["data"])
    tank500_rows = [
        row
        for row in static_rows
        if row.get("车型") == "坦克500" and row.get("问题") == "坦克500的续航"
    ]

    assert cleanup.removed_count == 1
    assert len(tank500_rows) == 1
    assert tank500_rows[0]["回答"] == "坦克500的WLTC纯电续航有110公里和201公里两个版本的，综合算下来的话有900公里和1096公里两个版本"
    assert list(tank500_rows[0]) == ["车型", "版本", "问题", "回答", "分类"]
