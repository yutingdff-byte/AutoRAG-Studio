from io import BytesIO
from types import SimpleNamespace

import pandas as pd
import pytest
from docx import Document

from knowledge.excel_restore import KnowledgeRestoreError, restore_excel
from knowledge.restore_manager import get_file_fingerprint, get_restore_cache_key, restore
from knowledge.word_fast_restore import restore_word_fast
from knowledge.word_restore import generate_material_rag, restore_word


class NamedBytesIO(BytesIO):
    def __init__(self, value: bytes, name: str):
        super().__init__(value)
        self.name = name


def make_excel_file(rows, name="历史知识.xlsx"):
    buffer = BytesIO()
    pd.DataFrame(rows).to_excel(buffer, index=False)
    buffer.seek(0)
    return NamedBytesIO(buffer.getvalue(), name)


def make_docx_file(paragraphs=None, tables=None, name="历史知识.docx"):
    document = Document()
    for paragraph in paragraphs or []:
        if isinstance(paragraph, tuple):
            text, style = paragraph
            document.add_paragraph(text, style=style)
        else:
            document.add_paragraph(paragraph)
    for table_rows in tables or []:
        table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for row_index, row in enumerate(table_rows):
            for col_index, value in enumerate(row):
                table.cell(row_index, col_index).text = value
    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return NamedBytesIO(buffer.getvalue(), name)


def test_restore_standard_excel_to_knowledge_items():
    file = make_excel_file(
        [
            {
                "车型": "理想MEGA",
                "版本": "全系",
                "问题": "MEGA多少钱？",
                "回答": "官方指导价52.98万元起。",
                "分类": "价格",
            }
        ]
    )

    items = restore_excel(file)

    assert len(items) == 1
    assert items[0].question == "MEGA多少钱？"
    assert items[0].model == "理想MEGA"
    assert items[0].trim == "全系"
    assert items[0].category == "价格"
    assert items[0].knowledge_type == "price"
    assert items[0].source_files == ["历史知识.xlsx"]


def test_restore_excel_rejects_non_standard_columns():
    file = make_excel_file(
        [
            {
                "问题": "多少钱？",
                "回答": "52.98万元起。",
            }
        ],
        name="非标准.xlsx",
    )

    with pytest.raises(KnowledgeRestoreError, match="不是标准知识库"):
        restore_excel(file)


def test_restore_word_reuses_existing_parser_and_agents(monkeypatch):
    calls = []
    logs = []

    def fake_parse(file):
        calls.append(("parse", file.name))
        return "历史Word资料"

    def fake_extract(material):
        calls.append(("facts", material))
        return {
            "facts": [
                {
                    "fact_id": "F001",
                    "content": "MEGA官方指导价52.98万元起。",
                }
            ]
        }

    def fake_rag(facts, progress_callback=None):
        calls.append(("rag", facts["facts"][0]["fact_id"]))
        if progress_callback:
            progress_callback("生成RAG", "批次成功", "Batch 1/1")
        return {
            "rag_knowledge": [
                {
                    "rag_id": "RAG-001",
                    "questions": ["MEGA多少钱？"],
                    "answer": "官方指导价52.98万元起。",
                    "category": "价格",
                    "knowledge_type": "dynamic",
                    "model": "理想MEGA",
                    "trim": "全系",
                    "fact_refs": ["F001"],
                }
            ]
        }

    monkeypatch.setattr("knowledge.word_restore.parse_document", fake_parse)
    monkeypatch.setattr("knowledge.word_restore.extract_material_facts", fake_extract)
    monkeypatch.setattr("knowledge.word_restore.generate_material_rag", fake_rag)

    items = restore_word(
        SimpleNamespace(name="产品FAQ.docx"),
        progress_callback=lambda stage, status, message: logs.append((stage, status, message)),
    )

    assert calls == [
        ("parse", "产品FAQ.docx"),
        ("facts", "历史Word资料"),
        ("rag", "F001"),
    ]
    assert len(items) == 1
    assert items[0].source_files == ["产品FAQ.docx"]
    assert items[0].fact_refs == ["F001"]
    assert ("深度提取Facts", "开始") in [(entry[0], entry[1]) for entry in logs]
    assert ("生成RAG", "开始") in [(entry[0], entry[1]) for entry in logs]
    assert logs[-1][0] == "完成"


def test_restore_manager_combines_success_and_failure():
    valid = make_excel_file(
        [
            {
                "车型": "理想MEGA",
                "版本": "全系",
                "问题": "MEGA多少钱？",
                "回答": "官方指导价52.98万元起。",
                "分类": "价格",
            }
        ],
        name="标准知识.xlsx",
    )
    invalid = SimpleNamespace(name="历史资料.pdf")

    result = restore([valid, invalid])

    assert result.restored_count == 1
    assert result.model_count == 1
    assert len(result.files) == 2
    assert result.files[0].success
    assert not result.files[1].success
    assert any("Restore" in entry for entry in result.files[1].logs)


def test_restore_manager_reports_word_stage_failure(monkeypatch):
    def fake_parse(file):
        return "历史Word资料"

    def fake_extract(material):
        print("JSON解析失败")
        print("模型返回不是合法JSON")
        return None

    monkeypatch.setattr("knowledge.word_restore.parse_document", fake_parse)
    monkeypatch.setattr("knowledge.word_restore.extract_material_facts", fake_extract)

    result = restore(SimpleNamespace(name="产品FAQ.docx"))

    assert result.restored_count == 0
    assert not result.files[0].success
    assert "知识恢复失败" in result.files[0].error
    assert any("读取文档｜成功" in entry for entry in result.files[0].logs)
    assert any("深度提取Facts｜开始" in entry for entry in result.files[0].logs)
    assert any("JSON解析失败" in entry for entry in result.files[0].logs)
    assert any("Restore｜失败" in entry for entry in result.files[0].logs)


def test_file_fingerprint_is_stable_for_same_uploaded_content():
    first = NamedBytesIO(b"same content", "history.docx")
    second = NamedBytesIO(b"same content", "history.docx")
    changed = NamedBytesIO(b"changed content", "history.docx")

    assert get_file_fingerprint(first) == get_file_fingerprint(second)
    assert get_file_fingerprint(first) != get_file_fingerprint(changed)


def test_fast_restore_word_table_qa():
    file = make_docx_file(
        tables=[
            [
                ["车型", "版本", "问题", "回答", "分类"],
                ["哈弗H6", "全系", "哈弗H6多少钱？", "官方指导价以当前政策为准。", "价格"],
            ]
        ]
    )

    result = restore_word_fast(file)

    assert len(result.items) == 1
    assert result.items[0].question == "哈弗H6多少钱？"
    assert result.items[0].model == "哈弗H6"
    assert result.items[0].trim == "全系"
    assert result.table_items == 1


def test_fast_restore_question_answer_paragraphs():
    file = make_docx_file(
        paragraphs=[
            "问题：坦克300适合越野吗？",
            "回答：坦克300更偏硬派越野定位，适合关注通过性的用户。",
        ]
    )

    result = restore_word_fast(file)

    assert len(result.items) == 1
    assert result.items[0].question == "坦克300适合越野吗？"
    assert "硬派越野" in result.items[0].answer


def test_fast_restore_heading_body():
    file = make_docx_file(
        paragraphs=[
            "质保政策",
            "整车质保以官方政策为准，基础售后保障可向门店确认。",
            "价格政策",
            "不同车型价格不同，建议按用户关注车型介绍。",
        ]
    )

    result = restore_word_fast(file)

    assert len(result.items) == 2
    assert result.heading_items == 2
    assert result.items[0].question == "质保政策"


def test_fast_restore_json_paragraphs_prioritize_json_fields():
    file = make_docx_file(
        paragraphs=[
            '{"车系":"25款新哈弗H5","问题":"25款新哈弗H5汽油领航的价格是多少？","答案":"价格是16.78万元，没有算上优惠哦。"}'
            '{"车系":"哈弗猛龙","问题":"哈弗猛龙26款燃油版的红包礼是什么？","答案":"资料中没有明确红包礼信息。"}',
        ]
    )

    result = restore_word_fast(file)

    assert len(result.items) == 2
    assert result.json_items == 2
    assert result.failed_items == 0
    assert result.confidence == 1
    assert result.items[0].model == "25款新哈弗H5"
    assert result.items[0].question == "25款新哈弗H5汽油领航的价格是多少？"
    assert not any("答案" in item.question for item in result.items)
    assert not any("{" in item.answer for item in result.items)


def test_fast_restore_json_in_table_cell():
    file = make_docx_file(
        tables=[
            [
                ['{"车型":"风骏5","问题":"风骏5是什么车型？","答案":"风骏5是燃油车型。"}'],
            ]
        ]
    )

    result = restore_word_fast(file)

    assert len(result.items) == 1
    assert result.items[0].model == "风骏5"
    assert result.items[0].answer == "风骏5是燃油车型。"


def test_fast_restore_rejects_json_leakage_in_question_answer():
    file = make_docx_file(
        paragraphs=[
            '"答案":"风骏5是燃油车型。"',
            '}{ "车系":"风骏5","问题":"风骏5是什么车型？","答案":"风骏5是燃油车型。" }',
        ]
    )

    result = restore_word_fast(file)

    assert len(result.items) == 1
    assert result.failed_items >= 1
    assert result.items[0].question == "风骏5是什么车型？"


def test_restore_word_fast_path_skips_agents(monkeypatch):
    file = make_docx_file(
        tables=[
            [
                ["问题", "回答", "分类"],
                ["欧拉好猫有什么权益？", "当前权益以官方政策资料为准。", "权益"],
            ]
        ]
    )

    def fail_extract(material):
        raise AssertionError("fast restore should not call Fact Agent")

    monkeypatch.setattr("knowledge.word_restore.extract_material_facts", fail_extract)

    items = restore_word(file)

    assert len(items) == 1
    assert items[0].question == "欧拉好猫有什么权益？"


def test_restore_word_unstructured_falls_back_to_agents(monkeypatch):
    file = make_docx_file(
        paragraphs=[
            "这是一段没有问答结构的车型资料，包含基础信息和政策描述。",
        ],
        name="非结构化.docx",
    )
    calls = []

    def fake_extract(material):
        calls.append("facts")
        return {"facts": [{"fact_id": "F001", "content": "资料内容"}]}

    def fake_rag(facts, progress_callback=None):
        calls.append("rag")
        return {
            "rag_knowledge": [
                {
                    "rag_id": "RAG-001",
                    "questions": ["有什么信息？"],
                    "answer": "资料中包含基础信息。",
                    "category": "配置",
                    "knowledge_type": "static",
                    "fact_refs": ["F001"],
                }
            ]
        }

    monkeypatch.setattr("knowledge.word_restore.extract_material_facts", fake_extract)
    monkeypatch.setattr("knowledge.word_restore.generate_material_rag", fake_rag)

    items = restore_word(file)

    assert calls == ["facts", "rag"]
    assert len(items) == 1


def test_restore_word_low_confidence_fast_result_falls_back_to_agents(monkeypatch):
    file = make_docx_file(
        paragraphs=[
            '"answer":"broken json"',
            '{"model":"Fengjun5","question":"What is Fengjun5?","answer":"Fengjun5 is a fuel vehicle."}',
        ],
        name="低可信.docx",
    )
    calls = []

    def fake_extract(material):
        calls.append(("facts", "broken json" in material))
        return {"facts": [{"fact_id": "F001", "content": "深度恢复事实"}]}

    def fake_rag(facts, progress_callback=None):
        calls.append(("rag", True))
        return {
            "rag_knowledge": [
                {
                    "rag_id": "RAG-001",
                    "questions": ["深度恢复问题？"],
                    "answer": "深度恢复回答。",
                    "category": "配置",
                    "knowledge_type": "static",
                    "fact_refs": ["F001"],
                }
            ]
        }

    monkeypatch.setattr("knowledge.word_restore.extract_material_facts", fake_extract)
    monkeypatch.setattr("knowledge.word_restore.generate_material_rag", fake_rag)

    items = restore_word(file)

    assert calls == [("facts", True), ("rag", True)]
    assert len(items) == 1
    assert items[0].question == "深度恢复问题？"


def test_restore_cache_key_includes_strategy_and_deep_flag(monkeypatch):
    file = NamedBytesIO(b"same content", "history.docx")
    monkeypatch.setenv("DEEPSEEK_MODEL", "model-a")

    fast_key = get_restore_cache_key(file, deep_restore=False)
    deep_key = get_restore_cache_key(file, deep_restore=True)

    assert fast_key != deep_key
    assert fast_key == get_restore_cache_key(NamedBytesIO(b"same content", "history.docx"))


def test_deep_rag_partial_failure_keeps_completed_batches(monkeypatch):
    facts = [
        {"fact_id": "F001", "category": "配置", "content": "配置事实"},
        {"fact_id": "F002", "category": "价格", "content": "价格事实"},
    ]
    logs = []

    monkeypatch.setattr("agents.rag_agent.load_prompt", lambda: "prompt")
    monkeypatch.setattr("agents.rag_agent.normalize_facts", lambda value: value["facts"] if isinstance(value, dict) else value)
    monkeypatch.setattr("agents.rag_agent.normalize_vehicle_fields", lambda fact: None)
    monkeypatch.setattr("agents.rag_agent.build_fact_index", lambda value: {fact["fact_id"]: fact for fact in value})
    monkeypatch.setattr("agents.rag_agent.group_facts", lambda value: [[value[0]], [value[1]]])
    monkeypatch.setattr("agents.rag_agent.merge_small_groups", lambda groups, max_size=30: groups)
    monkeypatch.setattr("agents.rag_agent.normalize_rag_metadata", lambda result, fact_index: result)
    monkeypatch.setattr("agents.rag_agent.dedupe_rag_items", lambda items: items)
    monkeypatch.setattr("agents.rag_agent.ensure_model_price_overview", lambda result, facts: result)

    def fake_generate_batch(batch, system_prompt):
        if batch[0]["fact_id"] == "F002":
            return None
        return {
            "rag_knowledge": [
                {
                    "rag_id": "RAG-001",
                    "questions": ["配置怎么样？"],
                    "answer": "配置满足日常使用。",
                    "category": "配置",
                    "knowledge_type": "static",
                    "fact_refs": ["F001"],
                }
            ],
            "info_gaps": [],
            "confirm_items": [],
        }

    monkeypatch.setattr("agents.rag_agent.generate_batch", fake_generate_batch)
    monkeypatch.setattr("utils.rag_quality.apply_export_quality_gate", lambda result: result)

    result = generate_material_rag(
        {"facts": facts},
        progress_callback=lambda stage, status, message: logs.append((stage, status, message)),
    )

    assert len(result["rag_knowledge"]) == 1
    assert result["restore_failed_batches"] == [2]
    assert any(status == "批次失败" for _, status, _ in logs)
