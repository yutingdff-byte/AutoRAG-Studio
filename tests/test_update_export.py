from io import BytesIO

from openpyxl import load_workbook

from knowledge.models import KnowledgeItem
from pages.update_page import _build_update_export_files


def item(knowledge_id: str, knowledge_type: str, category: str) -> KnowledgeItem:
    return KnowledgeItem(
        knowledge_id=knowledge_id,
        question="测试问题？",
        answer="测试回答。",
        category=category,
        brand="测试品牌",
        model="测试车型",
        trim="全系",
        knowledge_type=knowledge_type,
        metadata={"diff_id": "DIFF-001", "review_decision": "ACCEPT_NEW"},
    )


def read_rows(data: bytes):
    wb = load_workbook(BytesIO(data))
    ws = wb.active
    return [tuple(cell.value for cell in row) for row in ws.iter_rows()]


def test_update_export_reuses_generate_excel_schema_and_splits_files():
    files = _build_update_export_files(
        [
            item("K1", "product", "配置"),
            item("K2", "price", "价格"),
        ]
    )

    assert set(files) == {"static", "dynamic"}
    static_rows = read_rows(files["static"]["data"])
    dynamic_rows = read_rows(files["dynamic"]["data"])

    assert static_rows[0] == ("车型", "版本", "问题", "回答", "分类")
    assert dynamic_rows[0] == ("车型", "版本", "问题", "回答", "分类")
    assert len(static_rows) == 2
    assert len(dynamic_rows) == 2
    assert "review_decision" not in static_rows[0]
    assert "diff_id" not in dynamic_rows[0]
    assert files["static"]["data"]
    assert files["dynamic"]["data"]

