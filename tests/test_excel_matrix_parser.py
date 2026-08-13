from io import BytesIO

from openpyxl import Workbook

from parser.excel_parser import (
    COLUMN_ORIENTED_VEHICLE_MATRIX,
    ROW_ORIENTED_TABLE,
    _detect_excel_orientation,
    _read_sheet_rows,
    parse_excel,
)


def _workbook_bytes(workbook: Workbook) -> bytes:
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def test_row_oriented_table_keeps_existing_orientation():
    workbook = Workbook()
    ws = workbook.active
    ws.append(["车型", "价格"])
    ws.append(["A车型", "10万"])
    ws.append(["B车型", "20万"])

    rows = _read_sheet_rows(ws)

    assert _detect_excel_orientation(rows) == ROW_ORIENTED_TABLE

    parsed = parse_excel(_workbook_bytes(workbook))

    assert "字段：车型 | 价格" in parsed
    assert "记录1：" in parsed
    assert "车型: A车型" in parsed
    assert "价格: 10万" in parsed


def test_column_oriented_vehicle_matrix_transposes_columns_to_records():
    workbook = Workbook()
    ws = workbook.active
    ws.append(["车型", "车系A", "车系B"])
    ws.append(["品牌", "品牌A", "品牌B"])
    ws.append(["车系名称", "A车系", "B车系"])
    ws.append(["车型名称", "Pro", "Max"])
    ws.append(["配置版本等级", "A-Pro", "B-Max"])
    ws.append(["厂商指导价", 100000, 200000])
    ws.append(["纯电续航里程", "100km", "200km"])

    rows = _read_sheet_rows(ws)

    assert _detect_excel_orientation(rows) == COLUMN_ORIENTED_VEHICLE_MATRIX

    parsed = parse_excel(_workbook_bytes(workbook))

    assert "【表格结构：横向车型矩阵】" in parsed
    assert "【车型记录1】" in parsed
    assert "车型: 车系A" in parsed
    assert "车系名称: A车系" in parsed
    assert "车型名称: Pro" in parsed
    assert "厂商指导价: 100000" in parsed
    assert "【车型记录2】" in parsed
    assert "车型: 车系B" in parsed
    assert "车型名称: Max" in parsed
    assert "厂商指导价: 200000" in parsed


def test_merged_vehicle_header_is_inherited_by_each_version_column():
    workbook = Workbook()
    ws = workbook.active
    ws.merge_cells("B1:D1")
    ws["A1"] = "车型"
    ws["B1"] = "车系A"
    ws.append(["品牌", "品牌A", "品牌A", "品牌A"])
    ws.append(["车系名称", "A车系", "A车系", "A车系"])
    ws.append(["车型名称", "Pro", "Max", "Ultra"])
    ws.append(["配置版本等级", "低配", "中配", "高配"])
    ws.append(["厂商指导价", 100000, 120000, 150000])

    parsed = parse_excel(_workbook_bytes(workbook))

    assert parsed.count("车型: 车系A") == 3
    assert "车型名称: Pro" in parsed
    assert "车型名称: Max" in parsed
    assert "车型名称: Ultra" in parsed


def test_empty_formatted_tail_columns_do_not_create_vehicle_records():
    workbook = Workbook()
    ws = workbook.active
    ws.append(["车型", "A", "B", "", ""])
    ws.append(["品牌", "品牌A", "品牌B", "", ""])
    ws.append(["车系名称", "A车系", "B车系", "", ""])
    ws.append(["车型名称", "Pro", "Max", "", ""])
    ws.append(["配置版本等级", "低配", "高配", "", ""])
    ws.append(["厂商指导价", 100000, 200000, "", ""])
    ws["H20"] = ""

    parsed = parse_excel(_workbook_bytes(workbook))

    assert parsed.count("【车型记录") == 2
    assert "字段8" not in parsed


def test_dynamic_policy_fields_remain_bound_to_their_vehicle_column():
    workbook = Workbook()
    ws = workbook.active
    ws.append(["车型", "A", "B"])
    ws.append(["品牌", "品牌A", "品牌B"])
    ws.append(["车系名称", "A车系", "B车系"])
    ws.append(["车型名称", "Pro", "Max"])
    ws.append(["配置版本等级", "低配", "高配"])
    ws.append(["厂商指导价", 100000, 200000])
    ws.append(["购车权益-购车补贴政策", "现金优惠1万", "现金优惠2万"])
    ws.append(["购车权益-厂家置换补贴", "置换补贴3000", "置换补贴5000"])
    ws.append(["购车权益-金融优惠", "2年0息", "3年0息"])
    ws.append(["购车权益-其他权益", "送保养", "送流量"])

    parsed = parse_excel(_workbook_bytes(workbook))
    record_1 = parsed.split("【车型记录1】", 1)[1].split("【车型记录2】", 1)[0]
    record_2 = parsed.split("【车型记录2】", 1)[1]

    assert "车型名称: Pro" in record_1
    assert "购车权益-购车补贴政策: 现金优惠1万" in record_1
    assert "购车权益-厂家置换补贴: 置换补贴3000" in record_1
    assert "购车权益-金融优惠: 2年0息" in record_1
    assert "现金优惠2万" not in record_1

    assert "车型名称: Max" in record_2
    assert "购车权益-购车补贴政策: 现金优惠2万" in record_2
    assert "购车权益-厂家置换补贴: 置换补贴5000" in record_2
    assert "购车权益-金融优惠: 3年0息" in record_2
    assert "现金优惠1万" not in record_2


def test_unknown_sparse_structure_falls_back_without_crashing():
    workbook = Workbook()
    ws = workbook.active
    ws.append(["说明", "这是一段说明"])
    ws.append(["备注"])

    parsed = parse_excel(_workbook_bytes(workbook))

    assert "【Sheet:" in parsed
    assert "记录1：" in parsed or "字段：" in parsed
