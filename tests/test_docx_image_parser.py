from io import BytesIO

import pytest
from docx import Document
from docx.shared import Inches
from PIL import Image

from parser import docx_parser


def _image_bytes(width=800, height=400, color=(240, 80, 80)) -> BytesIO:
    buffer = BytesIO()
    Image.new("RGB", (width, height), color).save(buffer, format="PNG")
    buffer.seek(0)
    return buffer


def _save_docx(tmp_path, name, paragraphs=None, images=None):
    doc = Document()
    for text in paragraphs or []:
        doc.add_paragraph(text)
    for image in images or []:
        doc.add_picture(image, width=Inches(4))
    path = tmp_path / name
    doc.save(path)
    return path


def test_text_dominant_word_does_not_call_vision(tmp_path, monkeypatch):
    path = _save_docx(
        tmp_path,
        "text.docx",
        paragraphs=["这是一段用于测试的车型政策资料。" * 20],
        images=[_image_bytes()],
    )
    calls = []

    def fake_parse_image(file):
        calls.append(file)
        return "不应调用"

    monkeypatch.setattr(docx_parser, "parse_image", fake_parse_image)

    result = docx_parser.parse_docx(str(path))

    assert "车型政策资料" in result
    assert "检测到该 Word 主要由图片组成" not in result
    assert calls == []


def test_image_dominant_word_uses_existing_image_parser(tmp_path, monkeypatch):
    path = _save_docx(
        tmp_path,
        "image.docx",
        paragraphs=["营销政策"],
        images=[_image_bytes()],
    )
    calls = []

    def fake_parse_image(file):
        calls.append(file.name)
        return "识别到价格政策和购车权益。"

    monkeypatch.setattr(docx_parser, "parse_image", fake_parse_image)

    result = docx_parser.parse_docx(str(path))

    assert "检测到该 Word 主要由图片组成，正在使用图片识别。" in result
    assert "识别到价格政策和购车权益。" in result
    assert len(calls) == 1
    assert calls[0].startswith("image_image_001")


def test_logo_only_word_does_not_trigger_image_dominant(tmp_path, monkeypatch):
    path = _save_docx(
        tmp_path,
        "logo.docx",
        paragraphs=["Logo"],
        images=[_image_bytes(width=80, height=40)],
    )
    calls = []

    def fake_parse_image(file):
        calls.append(file)
        return "不应调用"

    monkeypatch.setattr(docx_parser, "parse_image", fake_parse_image)

    result = docx_parser.parse_docx(str(path))

    assert "Logo" in result
    assert "检测到该 Word 主要由图片组成" not in result
    assert calls == []


def test_image_dominant_word_keeps_success_when_one_image_fails(tmp_path, monkeypatch):
    path = _save_docx(
        tmp_path,
        "partial.docx",
        paragraphs=[""],
        images=[_image_bytes(color=(10, 120, 240)), _image_bytes(color=(40, 180, 90))],
    )
    calls = []

    def fake_parse_image(file):
        calls.append(file.name)
        if file.name.endswith("_002.png"):
            raise RuntimeError("模拟图片识别失败")
        return "第一张图片识别成功。"

    monkeypatch.setattr(docx_parser, "parse_image", fake_parse_image)

    result = docx_parser.parse_docx(str(path))

    assert "第一张图片识别成功。" in result
    assert "图片 2/2 识别失败" in result
    assert "成功 1/2，失败 1/2" in result
    assert len(calls) == 2


def test_image_dominant_word_fails_when_all_images_fail_and_no_text(tmp_path, monkeypatch):
    path = _save_docx(
        tmp_path,
        "all_failed.docx",
        paragraphs=[""],
        images=[_image_bytes()],
    )

    def fake_parse_image(file):
        raise RuntimeError("模拟图片识别失败")

    monkeypatch.setattr(docx_parser, "parse_image", fake_parse_image)

    with pytest.raises(Exception, match="未能从该 Word 中识别到有效内容"):
        docx_parser.parse_docx(str(path))
